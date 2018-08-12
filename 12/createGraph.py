import openpyxl
from openpyxl import chart

wb = openpyxl.Workbook()
sheet = wb.active
for row in range(1, 11):
    sheet.cell(row, 1).value = row

ref_obj = chart.Reference(sheet, 1, 1, 1, 10)
series_obj = chart.Series(ref_obj, title='First series')
chart_obj = chart.BarChart()
chart_obj.append(series_obj)
chart_obj.y = 50
chart_obj.x = 100
chart_obj.w = 300
chart_obj.h = 200
sheet.add_chart(chart_obj)
wb.save('sampleChart.xlsx')
